#!/usr/bin/env python3
"""Minimal PPTX -> PNG preview renderer (layout checking only).

Why this exists: this cluster has no LibreOffice, and the Aspose route needs a
newer glibc than the system provides.  So we rasterise decks ourselves with
python-pptx + Pillow to eyeball spacing / overflow / overlaps before presenting.

It is an APPROXIMATION: fonts are substituted (Liberation Sans ~ Calibri),
autoshapes are drawn as plain/rounded rectangles, and block arrows are drawn
from polygons.  Good enough to judge layout, not pixel-accurate.

Usage:
    python3 pptx_preview.py deck.pptx [outdir] [dpi]
"""
import io
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL, MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

EMU_PER_IN = 914400.0
FONTS = {
    (False, False): '/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf',
    (True, False): '/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf',
    (False, True): '/usr/share/fonts/liberation-sans/LiberationSans-Italic.ttf',
    (True, True): '/usr/share/fonts/liberation-sans/LiberationSans-BoldItalic.ttf',
}
DEFAULT_FILL = (0x44, 0x72, 0xC4)   # theme accent fallback
DEFAULT_LINE = (0x7F, 0x7F, 0x7F)
PT_SCALE = 1.0                       # global font scale for the preview
_font_cache = {}


def font(size_pt, bold=False, italic=False):
    key = (round(size_pt, 1), bold, italic)
    if key not in _font_cache:
        px = max(6, int(round(size_pt * DPI / 72.0)))
        _font_cache[key] = ImageFont.truetype(FONTS[(bold, italic)], px)
    return _font_cache[key]


def px(v):
    return v / EMU_PER_IN * DPI


def _color(color_fmt, fallback):
    """Best-effort RGB from a ColourFormat."""
    try:
        if color_fmt.type == MSO_COLOR_TYPE.RGB:
            c = color_fmt.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    try:                                    # scheme colour (can't resolve) -> fallback
        if color_fmt.type == MSO_COLOR_TYPE.SCHEME:
            return fallback
    except Exception:
        pass
    return None


def fill_rgb(shape, fallback=None):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return _color(shape.fill.fore_color, DEFAULT_FILL)
    except Exception:
        pass
    return fallback


def line_rgb(shape, fallback=None):
    try:
        if shape.line.fill.type == MSO_FILL.SOLID:
            return _color(shape.line.color, DEFAULT_LINE)
    except Exception:
        pass
    return fallback


def _tokenize(runs):
    """runs -> list of explicit lines; each line is [(token, font, color)]."""
    lines = [[]]
    for txt, fnt, col in runs:
        for k, part in enumerate(txt.split('\n')):
            if k > 0:
                lines.append([])
            toks = part.split(' ')
            for j, tok in enumerate(toks):
                if tok:
                    lines[-1].append((tok, fnt, col))
                if j != len(toks) - 1:
                    lines[-1].append((' ', fnt, col))
    return lines


def wrap_runs(runs, max_w, draw):
    """Greedy word wrap, honouring explicit newlines. Returns list of lines."""
    out = []
    for toks in _tokenize(runs):
        cur, cur_w = [], 0.0
        for tok, fnt, col in toks:
            w = draw.textlength(tok, font=fnt)
            if tok != ' ' and cur and cur_w + w > max_w:
                out.append(cur)
                cur, cur_w = [], 0.0
            if not cur and tok == ' ':
                continue                      # no leading spaces on a wrapped line
            cur.append((tok, fnt, col))
            cur_w += w
        out.append(cur)
    return out or [[]]


def draw_textframe(draw, tf, x, y, w, h, default_size=18.0):
    """Render a text frame inside the given pixel box."""
    ml = px(tf.margin_left or 0)
    mr = px(tf.margin_right or 0)
    mt = px(tf.margin_top or 0)
    mb = px(tf.margin_bottom or 0)
    inner_w = max(8.0, w - ml - mr)

    blocks = []                                    # (lines, line_h, align, space_after)
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            size = r.font.size.pt if r.font.size is not None else default_size
            fnt = font(size * PT_SCALE, bool(r.font.bold), bool(r.font.italic))
            col = _color(r.font.color, (0x33, 0x33, 0x33)) or (0x33, 0x33, 0x33)
            runs.append((r.text, fnt, col))
        if not runs:
            blocks.append(([], 0, PP_ALIGN.LEFT, px(p.space_after or 0)))
            continue
        lines = wrap_runs(runs, inner_w, draw)
        fnts = [f for _, f, _ in runs]
        lh = max(f.size for f in fnts) * 1.22 * (p.line_spacing or 1.0)
        blocks.append((lines, lh, p.alignment, px(p.space_after or 0)))

    total_h = sum(len(ln) * lh + sa for ln, lh, _, sa in blocks)
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y + mt + max(0.0, (h - mt - mb - total_h) / 2.0)
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - total_h
    else:
        cy = y + mt

    for lines, lh, align, sa in blocks:
        for line in lines:
            lw = sum(draw.textlength(t, font=f) for t, f, _ in line)
            if align == PP_ALIGN.CENTER:
                cx = x + ml + (inner_w - lw) / 2.0
            elif align == PP_ALIGN.RIGHT:
                cx = x + w - mr - lw
            else:
                cx = x + ml
            for t, f, col in line:
                draw.text((cx, cy), t, font=f, fill=col)
                cx += draw.textlength(t, font=f)
            cy += lh
        cy += sa


def draw_arrow(draw, kind, x, y, w, h, color):
    """Block arrow as a polygon (approximation)."""
    if w <= 0 or h <= 0:
        return
    if kind in ('right', 'left'):
        head = min(w * 0.55, h)
        body = h * 0.42
        cy = y + h / 2.0
        if kind == 'right':
            pts = [(x, cy - body / 2), (x + w - head, cy - body / 2), (x + w - head, y),
                   (x + w, cy), (x + w - head, y + h), (x + w - head, cy + body / 2),
                   (x, cy + body / 2)]
        else:
            pts = [(x + w, cy - body / 2), (x + head, cy - body / 2), (x + head, y),
                   (x, cy), (x + head, y + h), (x + head, cy + body / 2), (x + w, cy + body / 2)]
    else:
        head = min(h * 0.55, w)
        body = w * 0.42
        cx = x + w / 2.0
        if kind == 'down':
            pts = [(cx - body / 2, y), (cx - body / 2, y + h - head), (x, y + h - head),
                   (cx, y + h), (x + w, y + h - head), (cx + body / 2, y + h - head),
                   (cx + body / 2, y)]
        else:
            pts = [(cx - body / 2, y + h), (cx - body / 2, y + head), (x, y + head),
                   (cx, y), (x + w, y + head), (cx + body / 2, y + head), (cx + body / 2, y + h)]
    draw.polygon(pts, fill=color)


ARROW_KINDS = {'RIGHT_ARROW': 'right', 'LEFT_ARROW': 'left', 'DOWN_ARROW': 'down', 'UP_ARROW': 'up'}


def draw_shape(draw, sh, img):
    try:
        x, y = px(sh.left or 0), px(sh.top or 0)
        w, h = px(sh.width or 0), px(sh.height or 0)
    except Exception:
        return
    if w <= 0 or h <= 0:
        return

    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for sub in sh.shapes:
            draw_shape(draw, sub, img)
        return

    if st == MSO_SHAPE_TYPE.PICTURE:
        try:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA')
            pic = pic.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
            img.alpha_composite(pic, (int(x), int(y)))
        except Exception as e:
            print('  [warn] picture failed:', e)
        return

    if getattr(sh, 'has_table', False) and sh.has_table:
        tbl = sh.table
        cy = y
        for ri, row in enumerate(tbl.rows):
            rh = px(row.height or 0)
            if rh <= 0:
                rh = h / max(1, len(tbl.rows))
            cx = x
            for ci, cell in enumerate(row.cells):
                cw = px(tbl.columns[ci].width or 0) if ci < len(tbl.columns) else w / len(row.cells)
                bg = fill_rgb(cell, (0xFF, 0xFF, 0xFF))
                if bg:
                    draw.rectangle([cx, cy, cx + cw, cy + rh], fill=bg)
                draw.rectangle([cx, cy, cx + cw, cy + rh], outline=(0xBF, 0xBF, 0xBF))
                draw_textframe(draw, cell.text_frame, cx, cy, cw, rh, default_size=14)
                cx += cw
            cy += rh
        return

    # generic autoshape / connector / textbox
    name = ''
    try:
        name = (sh.auto_shape_type and sh.auto_shape_type.name) or ''
    except Exception:
        name = ''
    if name in ARROW_KINDS:
        col = fill_rgb(sh, DEFAULT_FILL) or DEFAULT_FILL
        draw_arrow(draw, ARROW_KINDS[name], x, y, w, h, col)
    elif st == MSO_SHAPE_TYPE.LINE or name in ('LINE', 'STRAIGHT_CONNECTOR_1'):
        col = line_rgb(sh, (0x88, 0x88, 0x88)) or (0x88, 0x88, 0x88)
        draw.line([x, y, x + w, y + h], fill=col, width=2)
    else:
        fill = fill_rgb(sh)
        line = line_rgb(sh)
        if fill is None and line is None:
            fill = None                                   # pure textbox
        radius = 0.0
        if name.startswith('ROUNDED'):
            radius = min(w, h) * 0.18
        if fill is not None or line is not None:
            if radius > 0:
                draw.rounded_rectangle([x, y, x + w, y + h], radius=radius,
                                       fill=fill, outline=line, width=2)
            else:
                draw.rectangle([x, y, x + w, y + h], fill=fill, outline=line, width=2)

    if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
        txt = sh.text_frame.text.strip()
        if txt:
            draw_textframe(draw, sh.text_frame, x, y, w, h, default_size=14)


def render(path, outdir, dpi=110, only=None):
    global DPI, PT_SCALE
    DPI = dpi
    PT_SCALE = 1.0
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(path)
    W = int(round(px(prs.slide_width)))
    H = int(round(px(prs.slide_height)))
    outs = []
    for i, slide in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        img = Image.new('RGBA', (W, H), (0xFF, 0xFF, 0xFF, 0xFF))
        draw = ImageDraw.Draw(img)
        for sh in slide.shapes:
            draw_shape(draw, sh, img)
        out = os.path.join(outdir, 'slide%02d.png' % i)
        img.convert('RGB').save(out)
        outs.append(out)
    print('[ok] rendered %d slide(s) -> %s  (%dx%d)' % (len(outs), outdir, W, H))
    return outs


if __name__ == '__main__':
    src = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else '/tmp/pptx_preview'
    dpi = int(sys.argv[3]) if len(sys.argv) > 3 else 110
    render(src, outdir, dpi)
